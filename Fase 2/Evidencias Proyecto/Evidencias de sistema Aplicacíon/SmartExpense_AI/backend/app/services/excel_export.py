from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import Receipt, User


HEADERS = [
    "ID",
    "Archivo",
    "Estado",
    "Tipo",
    "Comercio",
    "RUT emisor",
    "RUT receptor",
    "Folio",
    "Fecha",
    "Neto",
    "IVA",
    "Total",
    "Confianza OCR",
    "Daño estimado",
    "Notas de reconstrucción",
]


def receipts_to_excel(user: User, receipts: list[Receipt]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comprobantes"

    title_font = Font(name="Calibri", size=16, bold=True, color="1F3D2B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3D2B")
    alt_fill = PatternFill("solid", fgColor="F4EFE6")
    thin = Border(
        left=Side(style="thin", color="D9D1C4"),
        right=Side(style="thin", color="D9D1C4"),
        top=Side(style="thin", color="D9D1C4"),
        bottom=Side(style="thin", color="D9D1C4"),
    )

    ws.merge_cells("A1:O1")
    ws["A1"] = f"CAPTON — reporte contable · {user.company_name}"
    ws["A1"].font = title_font
    ws.merge_cells("A2:O2")
    ws["A2"] = f"Generado para {user.full_name} ({user.email}) · datos procesados on-premise"
    ws["A2"].font = Font(italic=True, color="6B6258")

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    total_sum = 0.0
    for idx, rec in enumerate(receipts):
        row = 5 + idx
        values = [
            rec.id,
            rec.original_filename,
            rec.status,
            rec.document_type,
            rec.merchant,
            rec.rut_emisor,
            rec.rut_receptor,
            rec.folio,
            rec.issued_at,
            rec.net_amount,
            rec.tax_amount,
            rec.total_amount,
            rec.confidence,
            rec.damage_score,
            rec.reconstruction_notes,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin
            if idx % 2 == 1:
                cell.fill = alt_fill
        if rec.total_amount:
            total_sum += rec.total_amount

    summary_row = 6 + len(receipts)
    ws.cell(row=summary_row, column=11, value="TOTAL").font = Font(bold=True)
    ws.cell(row=summary_row, column=12, value=round(total_sum, 2)).font = Font(bold=True)

    widths = [8, 28, 12, 20, 28, 16, 16, 14, 14, 12, 12, 12, 14, 14, 48]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.auto_filter.ref = f"A4:O{4 + len(receipts)}"
    ws.freeze_panes = "A5"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
